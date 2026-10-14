import os
import re
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def hex_to_rgb(hex_str, default=(0x00, 0x5B, 0xAA)):
    if not hex_str:
        return default
    cleaned = hex_str.lstrip('#').strip()
    if len(cleaned) == 6:
        try:
            return tuple(int(cleaned[i:i+2], 16) for i in (0, 2, 4))
        except ValueError:
            return default
    return default

class BuildDocs:
    def __init__(self, doc_dir, output_filename=None, project_name=None, doc_title=None, brand_color=None):
        self.doc_dir = os.path.abspath(doc_dir)
        self.project_name = project_name or os.environ.get("PROJECT_NAME", "Project")
        self.doc_title = doc_title or os.environ.get("DOC_TITLE", "Technical & Architectural Documentation")
        
        env_filename = os.environ.get("OUTPUT_FILENAME")
        self.output_filename = output_filename or env_filename or "Technical_Documentation"
        
        color_hex = brand_color or os.environ.get("BRAND_COLOR", "#005BAA")
        self.brand_color_hex = color_hex.lstrip('#')
        self.brand_rgb = hex_to_rgb(self.brand_color_hex)

    def set_cell_background(self, cell, hex_color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tc_pr.append(shd)

    def set_cell_margins(self, cell, top=100, bottom=100, left=150, right=150):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tc_mar.append(node)
        tc_pr.append(tc_mar)

    def compile_modular_markdown(self, target_doc_dir=None):
        doc_dir = target_doc_dir or self.doc_dir
        topics_dir = os.path.join(doc_dir, 'topics') if os.path.exists(os.path.join(doc_dir, 'topics')) else os.path.join(doc_dir, 'topicos')
        us_dir = os.path.join(doc_dir, 'user_stories') if os.path.exists(os.path.join(doc_dir, 'user_stories')) else os.path.join(doc_dir, 'historico_us')

        header = (
            f"# {self.doc_title} - {self.project_name}\n\n"
            f"> *Note: This document is automatically and continuously compiled from modular documentation source files.*\n\n"
            "---\n\n"
        )

        content_parts = [header]

        topic_files = sorted(glob.glob(os.path.join(topics_dir, '[0-9]_*.md')))
        for path in topic_files:
            with open(path, 'r', encoding='utf-8') as f:
                content_parts.append(f.read().strip() + "\n\n---\n\n")

        us_files = sorted(glob.glob(os.path.join(us_dir, 'US-*.md')))
        
        topico_us = "## Topic 5: User Stories & Feature History\n\n"
        topico_us += "### 5.1 Traceability Matrix\n\n"
        topico_us += "| US ID | Title / Feature | Module | Status | Completion Date |\n"
        topico_us += "|---|---|---|---|---|\n"

        us_details = []
        if us_files:
            for us_path in us_files:
                with open(us_path, 'r', encoding='utf-8') as f:
                    us_content = f.read().strip()
                    us_details.append(us_content)

                    us_id_match = re.search(r'###\s*(US-\d+|US-[A-Za-z0-9\-]+):\s*(.*)', us_content)
                    status_match = re.search(r'-\s*\*\*Status:\*\*\s*(.*)', us_content, re.IGNORECASE)
                    date_match = re.search(r'-\s*\*\*(?:Completion Date|Data de Conclusão):\*\*\s*(.*)', us_content, re.IGNORECASE)
                    modulo_match = re.search(r'-\s*\*\*(?:Module|Módulo):\*\*\s*(.*)', us_content, re.IGNORECASE)

                    us_id = us_id_match.group(1) if us_id_match else "US"
                    title = us_id_match.group(2) if us_id_match else os.path.basename(us_path)
                    status = status_match.group(1) if status_match else "Completed"
                    date_val = date_match.group(1) if date_match else "-"
                    modulo_val = modulo_match.group(1) if modulo_match else "-"

                    topico_us += f"| {us_id} | {title} | {modulo_val} | {status} | {date_val} |\n"
        else:
            topico_us += "| - | No user stories recorded | - | - | - |\n"

        topico_us += "\n### 5.2 Detailed History by User Story / Bug Fix\n\n"
        if us_details:
            topico_us += "\n\n".join(us_details) + "\n"
        else:
            topico_us += "*[No user stories recorded yet in user_stories/]*\n"

        content_parts.append(topico_us)

        master_md_content = "".join(content_parts)
        master_md_filename = f"{self.output_filename}.md"
        master_md_path = os.path.join(doc_dir, master_md_filename)
        
        with open(master_md_path, 'w', encoding='utf-8') as f:
            f.write(master_md_content)

        return master_md_path

    def convert_md_to_docx(self, md_path, docx_path):
        if not os.path.exists(md_path):
            print(f"[ERROR] Markdown file '{md_path}' was not found.")
            return False

        with open(md_path, 'r', encoding='utf-8') as f:
            lines = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', line) for line in f.readlines()]

        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Arial'
        normal_style.font.size = Pt(10.5)
        normal_style.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)

        brand_rgb = RGBColor(*self.brand_rgb)
        secondary_rgb = RGBColor(max(0, self.brand_rgb[0] - 20), max(0, self.brand_rgb[1] - 20), max(0, self.brand_rgb[2] - 20))

        in_code_block = False
        code_lines = []
        in_table = False
        table_rows = []

        def flush_code_block():
            nonlocal in_code_block, code_lines
            if code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1F, 0x24, 0x30)
                code_lines = []
            in_code_block = False

        def flush_table():
            nonlocal in_table, table_rows
            if not table_rows:
                in_table = False
                return

            parsed_rows = []
            for r in table_rows:
                if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', r):
                    continue
                cells = [c.strip() for c in r.strip().strip('|').split('|')]
                if any(cells):
                    parsed_rows.append(cells)

            if parsed_rows:
                col_count = max(len(row) for row in parsed_rows)
                table = doc.add_table(rows=len(parsed_rows), cols=col_count)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                for r_idx, row_data in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(col_count):
                        cell = table.cell(r_idx, c_idx)
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell.text = val
                        self.set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                        
                        if is_header:
                            self.set_cell_background(cell, self.brand_color_hex)
                            for p in cell.paragraphs:
                                p.paragraph_format.space_before = Pt(2)
                                p.paragraph_format.space_after = Pt(2)
                                for run in p.runs:
                                    run.font.name = 'Arial'
                                    run.font.bold = True
                                    run.font.size = Pt(10)
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            bg_color = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
                            self.set_cell_background(cell, bg_color)
                            for p in cell.paragraphs:
                                p.paragraph_format.space_before = Pt(2)
                                p.paragraph_format.space_after = Pt(2)
                                for run in p.runs:
                                    run.font.name = 'Arial'
                                    run.font.size = Pt(9.5)
                                    run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)

            table_rows = []
            in_table = False

        for line in lines:
            raw_line = line.rstrip('\r\n')

            if raw_line.startswith('```'):
                if in_code_block:
                    flush_code_block()
                else:
                    if in_table:
                        flush_table()
                    in_code_block = True
                    code_lines = []
                continue

            if in_code_block:
                code_lines.append(raw_line)
                continue

            if raw_line.strip().startswith('|') and raw_line.strip().endswith('|'):
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(raw_line)
                continue
            elif in_table:
                flush_table()

            if not raw_line.strip():
                continue

            if raw_line.startswith('# '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(10)
                run = p.add_run(raw_line[2:].strip())
                run.font.name = 'Arial'
                run.font.size = Pt(20)
                run.font.bold = True
                run.font.color.rgb = brand_rgb

            elif raw_line.startswith('## '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(raw_line[3:].strip())
                run.font.name = 'Arial'
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = brand_rgb

            elif raw_line.startswith('### '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(raw_line[4:].strip())
                run.font.name = 'Arial'
                run.font.size = Pt(12.5)
                run.font.bold = True
                run.font.color.rgb = secondary_rgb

            elif raw_line.startswith('#### '):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(raw_line[5:].strip())
                run.font.name = 'Arial'
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)

            elif raw_line.startswith('- ') or raw_line.startswith('* '):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                content = raw_line[2:].strip()
                parts = re.split(r'(\*\*.*?\*\*)', content)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

            elif re.match(r'^\d+\.\s', raw_line):
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                content = re.sub(r'^\d+\.\s*', '', raw_line).strip()
                parts = re.split(r'(\*\*.*?\*\*)', content)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

            elif raw_line.startswith('> '):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(raw_line[2:].strip())
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)
                parts = re.split(r'(\*\*.*?\*\*)', raw_line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

        if in_table:
            flush_table()
        if in_code_block:
            flush_code_block()

        try:
            doc.save(docx_path)
            print(f"[OK] DOCX document successfully generated at: {docx_path}")
            return True
        except PermissionError:
            print(f"[WARNING] Could not save '{docx_path}' because it is open in another program (e.g., Microsoft Word).")
            print("          Close the document and re-run to regenerate the .docx file.")
            return False

if __name__ == '__main__':
    import subprocess
    import sys
    main_script = os.path.join(os.path.dirname(__file__), 'main.py')
    subprocess.run([sys.executable, main_script] + sys.argv[1:])
